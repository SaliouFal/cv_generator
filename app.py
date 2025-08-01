import os, json, uuid, shutil, subprocess, base64, re, requests
from pathlib import Path
from flask import (Flask, render_template, request, redirect, url_for,
                   session, send_file, flash, jsonify)
import docx, openai
from threading import Thread  # Ajout pour l'asynchrone
import time
import re
from typing import Dict
import pandas as pd
# ─────────────── Configuration générale ────────────────────────────
BASE_DIR = Path(__file__).resolve().parent
openai.api_key = os.getenv("OPENAI_API_KEY")
OPENAI_MODEL = "o3"

app = Flask(__name__)
app.config.update(
    SECRET_KEY=os.getenv("FLASK_SECRET", "change-me"),
    UPLOAD_FOLDER=BASE_DIR / "static" / "uploads",
    OUTPUT_FOLDER=BASE_DIR / "generated",
)
(app.config["UPLOAD_FOLDER"]).mkdir(parents=True, exist_ok=True)
(app.config["OUTPUT_FOLDER"]).mkdir(parents=True, exist_ok=True)

# Stockage des tâches en cours
background_tasks = {}
task_results = {}





# …
ALLOWED_EMAILS = set()         # global

def load_allowed_emails(path: str):
    global ALLOWED_EMAILS
    df = pd.read_excel(path, dtype=str)       # une colonne "email"
    ALLOWED_EMAILS = {e.strip().lower()
                      for e in df['email'].dropna()}
    
load_allowed_emails(BASE_DIR / "allowed.xlsx")



from functools import wraps
from flask import abort

def login_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if 'user_email' not in session:
            # 401 → non authentifié - ou rediriger carrément vers /login
            return redirect(url_for('login', next=request.path))
        return f(*args, **kwargs)
    return wrapper



@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "GET":
        return render_template("login.html",
                               next = request.args.get('next','/'))
    
    # ----- POST -----
    email = request.form.get("email","").strip().lower()
    
    if email in ALLOWED_EMAILS:
        session.clear()                  # on remet tout à zéro
        session["user_email"] = email
        dest = request.form.get("next") or url_for("index")
        return redirect(dest)
    
    flash("Adresse non reconnue.", "danger")
    return redirect(url_for("login"))

@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))











with open(BASE_DIR / "config.json", encoding="utf-8") as fh:
    CV_TEMPLATES = json.load(fh)

def get_template_meta(tid: str):
    return next((t for t in CV_TEMPLATES if t["id"] == tid), None)

# ─────────────── Helpers LaTeX & placeholders ───────────────────────
PLACEHOLDER_PAT = re.compile(r'%%[^%]+%%')
LATEX_ESC = {
    '\\': r'\textbackslash{}','{': r'\{','}': r'\}','$': r'\$','&': r'\&',
    '%': r'\%','#': r'\#','_': r'\_','^': r'\^{}','~': r'\~{}'
}



def latex_escape(txt) -> str:
    """Convertit n'importe quelle valeur en str puis échappe LaTeX."""
    if txt is None:
        txt = ""
    elif not isinstance(txt, str):
        txt = str(txt)                 # ⬅️  caste dict, list, int, etc.
    for c, repl in LATEX_ESC.items():
        txt = txt.replace(c, repl)
    return txt

# ─── helpers -----------------------------------------------------------------
ADDR_RE = re.compile(r"(\d{4,5})\s+([A-Za-zÀ-ÖØ-öø-ÿ’' -]{2,})", re.U)

# utils_artifacts.py  (ou en haut de app.py)
ART_DIR = BASE_DIR / "artifacts"
ART_DIR.mkdir(exist_ok=True)

def save_text(txt: str) -> str:
    fid = uuid.uuid4().hex
    (ART_DIR / f"{fid}.tex").write_text(txt, encoding="utf-8")
    return fid

def load_text(fid: str) -> str:
    return (ART_DIR / f"{fid}.tex").read_text(encoding="utf-8")

def save_json(obj: dict) -> str:
    fid = uuid.uuid4().hex
    (ART_DIR / f"{fid}.json").write_text(
        json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")
    return fid

def load_json(fid: str) -> dict:
    return json.loads((ART_DIR / f"{fid}.json").read_text(encoding="utf-8"))


# utils_artifacts.py
def save_json(obj):
    fid = uuid.uuid4().hex
    path = BASE_DIR / "artifacts" / f"{fid}.json"
    path.write_text(json.dumps(obj, ensure_ascii=False), encoding="utf-8")
    return fid
def load_json(fid):
    return json.loads((BASE_DIR / "artifacts" / f"{fid}.json").read_text(encoding="utf-8"))



def _split_address(addr: str) -> tuple[str, str]:
    """
    Route de Cocoyer, 97190 Le Gosier  →  ('Route de Cocoyer', '97190 Le Gosier')
    4 place de la Mairie 75004 Paris   →  ('4 place de la Mairie', '75004 Paris')
    Si aucun code postal n’est trouvé : tout en ligne 1, ligne 2 vide.
    """
    if not addr:
        return "", ""

    addr = addr.strip().replace("  ", " ")

    m = ADDR_RE.search(addr)
    if not m:
        # pas de code postal identifié
        return addr, ""

    # partie avant le CP : rue ou quartier
    line1 = addr[: m.start()].rstrip(", ").strip()
    # partie CP+ville
    line2 = f"{m.group(1)} {m.group(2)}".strip()

    # Fallback : si la rue est vide, on met tout en ligne 1
    if not line1:
        return line2, ""
    return line1, line2

def _skill_label(raw) -> str:
    """Renvoie toujours le libellé texte d’une compétence."""
    if isinstance(raw, str):
        return raw
    if isinstance(raw, dict):
        return raw.get("label") or raw.get("name") or raw.get("skill") or ""
    return str(raw)

# ─── build_placeholders  (remplace entièrement l’ancienne) -------------------
def build_placeholders(cv: Dict) -> Dict:
    """Build a mapping of LaTeX placeholders → rendered LaTeX strings.

    The degree blocks are now rendered with the same visual style as the
    experience blocks (colorbox + minipage), ensuring consistent design.
    """

    photo_file = cv.get("photo", "")
    if photo_file:
        absolute = app.config["UPLOAD_FOLDER"] / photo_file
        if not absolute.is_file():  # ⬅️ test sur le chemin absolu
            app.logger.warning(f"Fichier photo introuvable : {absolute}")
            photo_file = ""

    contacts = cv.get("contacts", {})
    phone = contacts.get("phone") or cv.get("phone", "")
    email = contacts.get("email") or cv.get("email", "")
    linkedin = contacts.get("linkedin") or cv.get("linkedin_url", "")

    add1, add2 = _split_address(contacts.get("address", ""))
    linkedin_handle = linkedin.rsplit('/', 1)[-1].lstrip('@') if linkedin else ""

    # ── placeholders simples ───────────────────────────────────────────────
    ph: Dict[str, str] = {
        "%%FULL_NAME%%": cv.get("full_name", ""),
        "%%HEADLINE%%": cv.get("headline", ""),
        "%%PHONE%%": phone,
        "%%EMAIL%%": email,
        "%%LINKEDIN_URL%%": linkedin,
        "%%LINKEDIN_HANDLE%%": linkedin_handle,
        "%%ADDRESS_LINE1%%": add1,
        "%%ADDRESS_LINE2%%": add2,
        "%%SUMMARY%%": cv.get("summary", ""),
        "%%PHOTO_FILE%%": cv.get("photo", ""),
    }

    # ─── LANGUES ───────────────────────────────────────────────────────────
    langs = cv.get("languages", [])
    if langs:
        lang_items = "\n".join(
            rf"\item {latex_escape(l.get('name',''))} - "
            rf"\textcolor{{gray}}{{{latex_escape(l.get('level',''))}}}"
            for l in langs
        )
        ph["%%LANGUAGE_ITEMS%%"] = (
            r"\begin{itemize}[leftmargin=*]" + "\n" + lang_items + r"\end{itemize}"
        )
    else:
        ph["%%LANGUAGE_ITEMS%%"] = " "

    # ─── COMPÉTENCES (max 7 – un seul mot) ────────────────────────────────
    skills = [_skill_label(s) for s in cv.get("skills", [])][:7]
    if skills:
        skill_items = "\n".join(rf"\item {latex_escape(s)}" for s in skills)
        ph["%%SKILL_ITEMS%%"] = (
            r"\begin{itemize}[leftmargin=*]" + "\n" + skill_items + r"\end{itemize}"
        )
    else:
        ph["%%SKILL_ITEMS%%"] = " "

    # ─── Grille 3×4 de compétences cerclées : %%SKILL_CIRCLES%% ───────────
    skills = [_skill_label(s) for s in cv.get("skills", [])]
    rows = []
    for i, lab in enumerate(skills):
        if i % 4 == 0:
            rows.append([])
        rows[-1].append(r"\cicon " + latex_escape(lab))

    if rows and len(rows[-1]) < 4:
        rows[-1] += ["~"] * (4 - len(rows[-1]))

    skill_lines = [" & ".join(r) + r" \\" for r in rows]
    ph["%%SKILL_CIRCLES%%"] = (
        r"\begin{tabular}{@{}p{0.25\linewidth}p{0.18\linewidth}p{0.18\linewidth}p{0.18\linewidth}}"
        + "\n".join(skill_lines)
        + r"\end{tabular}"
    )

    # ─── Centres d’intérêt ────────────────────────────────────────────────
    hobbies = cv.get("hobbies", []) or ["Lecture", "Sport", "Musique", "Voyage"]
    cv["hobbies"] = hobbies  # assure la présence pour d’autres traitements éventuels

    hobby_items = (
        r"\begin{itemize}[leftmargin=*]" + "\n" +
        "\n".join(rf"\item {latex_escape(h)}" for h in hobbies) + "\n" +
        r"\end{itemize}"
    )
    ph["%%HOBBY_ITEMS%%"] = hobby_items

    # ─── Expériences ───────────────────────────────────────────────────────
    jobs = cv.get("jobs") or cv.get("experiences") or []
    exp_blocks = []
    for j in jobs:
        bullets = j.get("description") or j.get("bullets") or []
        bullet_items = " \\item ".join(latex_escape(b) for b in bullets)
        block = rf"""\colorbox{{maincolor}}{{%
  \begin{{minipage}}{{\linewidth}}
    \noindent
    \textbf{{{latex_escape(j.get('title',''))}}}\hfill {j.get('dates','') or ''}\\
    {latex_escape(j.get('company',''))}\\[-0.3em]
    \begin{{itemize}}[leftmargin=*]
      \item {bullet_items}
    \end{{itemize}}
  \end{{minipage}}}}"""
        exp_blocks.append(block)
    ph["%%EXPERIENCE_BLOCKS%%"] = ("\n\n\\vspace{3mm}\n\n").join(exp_blocks) or " "

    # ─── Diplômes (même style que les expériences) ─────────────────────────
    deg_blocks = []
    for d in cv.get("degrees", []):
        title = d.get("title") or d.get("degree") or ""
        school = d.get("school") or d.get("institution") or ""
        dates = d.get("dates") or d.get("year") or ""
        desc = d.get("description", "")

        # bullets -----------------------------------------------------------
        bullets = []
        if isinstance(desc, list):
            bullets = [b for b in desc if b.strip()]
        elif desc:
            bullets = [s.strip(" •-–;\n\r\t") for s in re.split(r"[•\-–;|\n\r]", desc) if s.strip()]

        bullet_items = " \\item ".join(latex_escape(b) for b in bullets) if bullets else ""

        block = rf"""\colorbox{{maincolor}}{{%
  \begin{{minipage}}{{\linewidth}}
    \noindent
    \textbf{{{latex_escape(title)}}}\hfill {latex_escape(dates)}\\
    {latex_escape(school)}\\[-0.3em]
    \begin{{itemize}}[leftmargin=*]
      \item {bullet_items}
    \end{{itemize}}
  \end{{minipage}}}}"""
        deg_blocks.append(block)

    ph["%%DEGREE_BLOCKS%%"] = ("\n\n\\vspace{3mm}\n\n").join(deg_blocks) or " "

    # ─── Escaping final (pas pour les blocs LaTeX) ─────────────────────────
    raw_tags = {
        "%%LANGUAGE_ITEMS%%",
        "%%SKILL_ITEMS%%",
        "%%LANGUAGE_TABLE%%",
        "%%SKILL_ROWS%%",
        "%%EXPERIENCE_BLOCKS%%",
        "%%DEGREE_BLOCKS%%",
        "%%HOBBY_ITEMS%%",
        "%%SKILL_CIRCLES%%",
    }

    clean = {}
    for tag, val in ph.items():
        clean[tag] = val if tag in raw_tags else latex_escape(val or "")
    return clean


def apply_placeholders(tex: str, ph: Dict) -> str:
    for tag, val in ph.items():
        tex = tex.replace(tag, val)
    return PLACEHOLDER_PAT.sub("", tex)

######################################################################


# ------------------------------------------------------------------
def gpt_diagnose(cv_json: dict) -> dict:
    """Retourne un diagnostic structurÉ du CV."""
    sys = (
        "Tu es un recruteur senior. On te fournit le contenu JSON d’un CV "
        "et tu dois identifier tous les problèmes de fond (informations manquantes, phrases vagues, etc.) mais pas d'analyse sur la photo ou des trucs inutils.\n"
        "**Important** : Il faut tout mettre en Français.\n"
        "Réponds en JSON avec :\n"
        "  score   – note globale sur 20,\n"
        "  issues  – liste d’objets {zone , problem, why, suggestion},\n"
        "  priorities – tableau de 3 à 6 actions classées par urgence."
    )

    rsp = openai.ChatCompletion.create(
        model           = OPENAI_MODEL,
        response_format = {"type": "json_object"},
        messages=[{"role": "system", "content": sys},
                  {"role": "user",   "content": json.dumps(cv_json, ensure_ascii=False)}]
    )
    return json.loads(rsp.choices[0].message.content)


@app.route("/cv_diagnostic")
def cv_diagnostic():
    if "cv_id" not in session:                 # sécurité
        return jsonify(error="no CV"), 400

    cv_json = load_json(session["cv_id"])
    try:
        diag = gpt_diagnose(cv_json)
        return jsonify(diag)
    except Exception as e:
        app.logger.error("Diagnostic error: %s", e)
        return jsonify(error=str(e)), 500



# Modifier la fonction diag_prepare
@app.route("/diag_prepare", methods=["POST"])
def diag_prepare():
    if "cv_id" not in session:
        return jsonify(error="no CV"), 400

    cv_json = load_json(session["cv_id"])
    task_id = "diag_" + str(uuid.uuid4())  # Préfixe pour identification

    Thread(target=async_gpt_diagnose,
           args=(cv_json, task_id),
           daemon=True).start()

    return jsonify(task_id=task_id)

# Modifier async_gpt_diagnose
def async_gpt_diagnose(cv_json: dict, task_id: str):
    try:
        result = gpt_diagnose(cv_json)
        task_results[task_id] = {"status": "completed", "result": result}
    except Exception as e:
        task_results[task_id] = {"status": "error", "message": str(e)}



@app.route("/get_diagnostic/<task_id>")
def get_diagnostic(task_id):
    result = task_results.get(task_id, {})
    return jsonify(result)







######################################################################
# ─────────────── GPT helpers (enrichissement / parsing) ─────────────
SCHEMA = json.loads((BASE_DIR / "schema.json").read_text())
def gpt_enrich(partial: dict) -> dict:
    sys = (
    # — Rôle général ----------------------------------------------------------
    "Tu es un expert RH et rédacteur de CV en Français.\n"
   
    # — Contraintes structurelles --------------------------------------------
    "• Complète ou extrait les informations en STRICT respect du schéma JSON fourni.\n"
    "• Aucune invention de dates, d’entreprises ou de diplômes inexistants.\n"

    # — Résumé professionnel --------------------------------------------------
    "• Champ « summary » : s’il existe, reformule-le en 3 – 4 phrases fluides ; "
    "s’il est absent rédige-le en 3 ou 4 phrases.\n"

    # — Expériences -----------------------------------------------------------
    "• experiences[*].bullets : reformule chaque point en 2 – 3 puces synthétiques "
    "(verbe d’action + résultat), sans rien inventer.\n"

    # — Formations ------------------------------------------------------------
    "• **degrees[*].description** :  "
    "- Si la description est absente, écris 2-3 puces génériques sur le contenu du cursus.\n"
    "- S’il est vide ou absent, rédige une courte description (2-3 puces génériques sur le contenu du cursus) "
    "à partir du titre du diplôme, sans inventer de données factices.\n"

    # — Compétences -----------------------------------------------------------
    "• skills : [] La liste des compétences clés chaque compétences. **Important** : Chaque compétence doit etre résumée en un seul mot seulement. \n"

    # — Langues ---------------------------------------------------------------
    "• languages : tableau [{name, level}] séparé des compétences.\n"

    # — Interets ---------------------------------------------------------------
    "• Hobbies : tableau [""] .\n"

    # — Sortie ----------------------------------------------------------------
    "Réponds UNIQUEMENT avec l’objet JSON demandé."
)

    msg = {"schema":SCHEMA,"partial":partial}
    out = openai.ChatCompletion.create(
        model=OPENAI_MODEL,
        response_format={"type":"json_object"},
        messages=[{"role":"system","content":sys},
                  {"role":"user","content":json.dumps(msg,ensure_ascii=False)}]
    ).choices[0].message.content
    
    return json.loads(out)

def gpt_parse_plaintext(txt:str) -> dict:
    sys = (
    # — Rôle général ----------------------------------------------------------
    "Tu es un expert RH et rédacteur de CV en français; \n"
    
    # — Contraintes structurelles --------------------------------------------
    "• Complète ou extrait les informations en STRICT respect du schéma JSON fourni.\n"
    "• Aucune invention de dates, d’entreprises ou de diplômes inexistants.\n"
    
    

    # — Résumé professionnel --------------------------------------------------
    "• Champ « summary » : s’il existe, reformule-le en 3 – 4 phrases fluides "
    "s’il est absent ou < 20 mots, rédige-le.\n"

    # — Expériences -----------------------------------------------------------
    "• experiences[*].bullets : reformule chaque point en 2 – 3 puces synthétiques "
    "(verbe d’action + résultat), sans rien inventer.\n"

    # — Formations ------------------------------------------------------------
    "• **degrees[*].description** :  "
    "- Si la description est absente, écris 2-3 puces génériques sur le contenu du cursus.\n"
    "- S’il est vide ou absent, rédige une courte description (2-3 puces génériques sur le contenu du cursus) "
    "à partir du titre du diplôme, sans inventer de données factices.\n"

    # — Compétences -----------------------------------------------------------
    "• skills : retourne **maximum 10** libellés, **un seul mot** chacun ; "
   

    # — Langues ---------------------------------------------------------------
    "• languages : tableau [{name, level}] séparé des compétences.\n"

    # — Interets ---------------------------------------------------------------
    "• Hobbies : tableau [""] mets les interets dedans séparés.\n"


    # — Sortie ----------------------------------------------------------------
    "Réponds UNIQUEMENT avec l’objet JSON demandé."
)
    msg = {"schema":SCHEMA,"text":txt[:6000]}
    out = openai.ChatCompletion.create(
        model=OPENAI_MODEL,
        response_format={"type":"json_object"},
        messages=[{"role":"system","content":sys},
                  {"role":"user","content":json.dumps(msg,ensure_ascii=False)}]
    ).choices[0].message.content
    
    return json.loads(out)



# ───────────────────────── gpt_refine_json ──────────────────────────
def gpt_refine_json(cv_json: dict) -> dict:
    """
    Passe-2 : relecture stylistique + resserrage (1 page).
    """
    sys = (
        # rôle et périmètre ────────────────────────────────────────────────
        "Tu es un expert RH/rédacteur CV. Tu reçois l'objet JSON complet d'un CV.\n"
        # corrections globales ─────────────────────────────────────────────
        "0. Regarde le headline et si c'est pas bien formulé tu le reformules de manière plus coherente et professionnelle"
        "1. Corrige toutes fautes d’orthographe, de grammaire et de typographie.\n"
        "2. Remplace les phrases ou mots vagues par des formulations précises et professionnelles.\n"
        "3. Uniformise les dates au format « MM/AAAA - MM/AAAA » (ou « depuis MM/AAAA »).\n"
        "4. Supprime toute redondance, incohérence ou information hors-sujet.\n"
      
        # contraintes structurelles ───────────────────────────────────────
        "6. Ne change en AUCUN CAS la structure : mêmes clés, mêmes niveaux.\n"
        "7. Si le cv Semble trop long et il risque de faire 2 pages tu devra essayer de resumer un peu les descriptions des formations et experiences pour que le cv tienne sur une seule page.\n"
        # réponse attendue ────────────────────────────────────────────────
        "Réponds STRICTEMENT avec l’objet JSON final, rien d’autre."
    )

    rsp = openai.ChatCompletion.create(
        model           = OPENAI_MODEL,
        response_format = {"type": "json_object"},
        messages=[{"role": "system", "content": sys},
                  {"role": "user",   "content": json.dumps(cv_json, ensure_ascii=False)}]
    )
    
    return json.loads(rsp.choices[0].message.content)




# ─────────────── Compilation LaTeX distante ────────────────────────
LATEX_API = "https://latex.ytotech.com/builds/sync"
def copy_template(meta):
    dst = app.config["OUTPUT_FOLDER"]/uuid.uuid4().hex
    shutil.copytree(Path(meta["latex_path"]).parent, dst)
    return dst

def collect_assets(build:Path):
    exts={".jpg",".jpeg",".png",".pdf",".eps",".svg"}
    return {p.relative_to(build).as_posix():p for p in build.rglob("*")
            if p.is_file() and p.suffix.lower() in exts}

def compile_pdf_remote(tex: str, build: Path, extra: dict | None = None) -> Path:
    resources = [{"main": True, "content": tex}]
    
    # Ajouter les ressources supplémentaires
    for rel, p in (extra or {}).items():
        if p.exists():
            resources.append({
                "path": rel,
                "file": base64.b64encode(p.read_bytes()).decode()
            })
        else:
            app.logger.warning(f"Fichier supplémentaire introuvable: {p}")

    try:
        r = requests.post(
            LATEX_API,
            json={"compiler": "xelatex", "resources": resources},
            timeout=90
        )
        
        if r.status_code >= 300 or not r.headers.get("content-type", "").startswith("application/pdf"):
            error_msg = f"LaTeX API ERROR {r.status_code}: {r.text[:500]}"
            app.logger.error(error_msg)
            
            # Sauvegarder le code LaTeX pour débogage
            tex_path = build / "error.tex"
            tex_path.write_text(tex, encoding="utf-8")
            app.logger.error(f"Code LaTeX sauvegardé à: {tex_path}")
            
            raise RuntimeError(error_msg)
        
        pdf = build / "main.pdf"
        pdf.write_bytes(r.content)
        return pdf
        
    except Exception as e:
        app.logger.exception("Erreur dans compile_pdf_remote")
        
        # Sauvegarder le code LaTeX pour débogage
        tex_path = build / "error.tex"
        tex_path.write_text(tex, encoding="utf-8")
        app.logger.error(f"Code LaTeX sauvegardé à: {tex_path}")
        
        raise
# ─────────────── Routes ────────────────────────────────────────────
@app.route("/")
@login_required
def index(): return render_template("index.html", templates=CV_TEMPLATES)

@app.route("/select/<template_id>")
def select_template(template_id):
    if not get_template_meta(template_id): return redirect(url_for("index"))
    session["template_id"]=template_id
    return redirect(url_for("choose_input"))

@app.route("/choose_input")
@login_required
def choose_input():
    if "template_id" not in session: return redirect(url_for("index"))
    return render_template("choose_input.html")

# ---------- 1. Formulaire -----------------------------------------
array_re = re.compile(r"(?P<fld>jobs|degrees)\[(?P<idx>\d+)]\[(?P<key>\w+)]")
def _first(v): return v[0] if isinstance(v,list) else v
def extract_arrays(form)->dict:
    buf={}
    for k,v in (form.items() if hasattr(form,"items") else form):
        m=array_re.fullmatch(k); 
        if not m: continue
        fld,idx,key=m.group("fld"),int(m.group("idx")),m.group("key")
        buf.setdefault((fld,idx),{})[key]=_first(v).strip()
    out={"jobs":[], "degrees":[], "languages":[], "skills":[]}
    for (fld,idx),row in sorted(buf.items(), key=lambda p:p[0][1]):
        out.setdefault(fld,[]).append(row)
    return out

@app.route("/minidata", methods=["GET","POST"])
@login_required
def minidata():
    if request.method=="GET": return render_template("minidata.html")
    session["raw_form"]=request.form.to_dict(flat=False)
    photo=request.files.get("photo")
    if photo and photo.filename:
        name=f"{uuid.uuid4().hex}{Path(photo.filename).suffix}"
        dst=app.config["UPLOAD_FOLDER"]/name; dst.parent.mkdir(parents=True,exist_ok=True)
        photo.save(dst); session["photo_file"]=name
    else: session["photo_file"]=""
    return redirect(url_for("generate_from_form"))


def async_gpt_enrich(partial: dict, task_id: str):
    """Version asynchrone de gpt_enrich"""
    try:
        result = gpt_enrich(partial)
        result = gpt_refine_json(result)  
        task_results[task_id] = {"status": "completed", "result": result}
    except Exception as e:
        task_results[task_id] = {"status": "error", "message": str(e)}

def async_gpt_parse_plaintext(txt: str, task_id: str):
    """Version asynchrone de gpt_parse_plaintext"""
    try:
        result = gpt_parse_plaintext(txt)
        result = gpt_refine_json(result)  
        task_results[task_id] = {"status": "completed", "result": result}
    except Exception as e:
        task_results[task_id] = {"status": "error", "message": str(e)}

def async_gpt_edit_json(cv_json: dict, instruction: str, task_id: str):
    """Version asynchrone de gpt_edit_json"""
    try:
        result = gpt_edit_json(cv_json, instruction)
        task_results[task_id] = {"status": "completed", "result": result}
    except Exception as e:
        task_results[task_id] = {"status": "error", "message": str(e)}

def async_gpt_refine(cv_json: dict, task_id: str):
    try:
        result = gpt_refine_json(cv_json)
        task_results[task_id] = {"status": "completed", "result": result}
    except Exception as e:
        task_results[task_id] = {"status": "error", "message": str(e)}


# =============================================================================
# Routes modifiées pour utiliser l'asynchrone
# =============================================================================
@app.route("/processing/<task_id>")
def processing(task_id):
    return render_template("processing.html", task_id=task_id)

@app.route("/generate_from_form")

def generate_from_form():
    task_id = str(uuid.uuid4())
    session["form_task_id"] = task_id
    
    # Préparation des données
    flat = {k: v[0] for k, v in session.pop("raw_form").items()}
    flat.update(extract_arrays(flat))
    flat["photo"] = session.pop("photo_file", "")
    
    # Lancement de la tâche en arrière-plan
    thread = Thread(target=async_gpt_enrich, args=(flat, task_id))
    thread.start()
    
    return render_template("processing.html", task_id=task_id)

@app.route("/import_cv", methods=["POST"])

def import_cv():
    try:
        # Vérification de la session
        if "template_id" not in session:
            return jsonify({"error": "Session expirée", "redirect": url_for('index')}), 401

        # Vérification du fichier
        if 'cvfile' not in request.files:
            return jsonify({"error": "Aucun fichier n'a été envoyé"}), 400
            
        up = request.files['cvfile']
        
        if not up or up.filename == '':
            return jsonify({"error": "Aucun fichier sélectionné"}), 400
        
        # Vérification de l'extension
        ext = Path(up.filename).suffix.lower()
        if ext not in ['.pdf', '.doc', '.docx']:
            return jsonify({"error": "Format de fichier non supporté"}), 400
        
        

        # Traitement du fichier
        tmp = app.config["UPLOAD_FOLDER"] / f"tmp_{uuid.uuid4().hex}{ext}"
        tmp.parent.mkdir(parents=True, exist_ok=True)
        up.save(tmp)
        
        photo = request.files.get('photo')
        photo_name = ""
        if photo and photo.filename:
            photo_name = f"{uuid.uuid4().hex}{Path(photo.filename).suffix}"
            photo.save(app.config["UPLOAD_FOLDER"] / photo_name)

        # Extraction du texte
        if ext == ".pdf":
            plain = subprocess.check_output(["pdftotext", "-layout", str(tmp), "-"], text=True)
        else:
            plain = "\n".join(p.text for p in docx.Document(str(tmp)).paragraphs)
        
        tmp.unlink(missing_ok=True)
        
        task_id = str(uuid.uuid4())
        session["import_task_id"] = task_id
        session["photo_name"] = photo_name
        
        # Lancement de la tâche en arrière-plan
        thread = Thread(target=async_gpt_parse_plaintext, args=(plain, task_id))
        thread.start()
        
        return jsonify({
            "status": "processing",
            "import_task_id": task_id,
            "redirect": url_for('processing', task_id=task_id)
        })
        
    except Exception as e:
        app.logger.error(f"Erreur dans import_cv: {str(e)}")
        return jsonify({"error": f"Erreur de traitement: {str(e)}"}), 500
    


@app.route("/check_task/<task_id>")
def check_task(task_id):
    result = task_results.get(task_id, {})
    
    # Si la tâche est terminée, on finalise
    if result.get("status") == "completed":
        # Pour /generate_from_form
        if task_id == session.get("form_task_id"):
            cv_json = result["result"]
            # print(cv_json)
            meta = get_template_meta(session["template_id"])
            tpl = Path(meta["latex_path"]).read_text(encoding="utf-8")
            tex = apply_placeholders(tpl, build_placeholders(cv_json))

            build = copy_template(meta)
            extra = collect_assets(build)
            if cv_json.get("photo"):
                extra[cv_json["photo"]] = app.config["UPLOAD_FOLDER"] / cv_json["photo"]
            pdf = compile_pdf_remote(tex, build, extra)

            session.update({
                "last_tex": tex,
                "pdf_filename": pdf.relative_to(app.config["OUTPUT_FOLDER"]).as_posix()
            })
            return jsonify(status="completed", redirect=url_for("preview"))
        
        # Pour /import_cv
        elif task_id == session.get("import_task_id"):
            cv_json = result["result"]
            cv_json["photo"] = session.get("photo_name", "")
            

            meta = get_template_meta(session["template_id"])
            tpl = Path(meta["latex_path"]).read_text(encoding="utf-8")
            tex = apply_placeholders(tpl, build_placeholders(cv_json))

            build = copy_template(meta)
            extra = collect_assets(build)
            if cv_json["photo"]:
                extra[cv_json["photo"]] = app.config["UPLOAD_FOLDER"] / cv_json["photo"]
            pdf = compile_pdf_remote(tex, build, extra)
            
            tex_id = save_text(tex)
            cv_id = save_json(cv_json)

            session.update({
                "tex_id": tex_id,
                "cv_id": cv_id,
                "pdf_filename": pdf.relative_to(app.config["OUTPUT_FOLDER"]).as_posix()
            })
            return jsonify(status="completed", redirect=url_for("preview"))
        
        # Pour /json_edit_apply
        elif task_id == session.get("edit_task_id"):
            updated_json = result["result"]
            
            meta = get_template_meta(session["template_id"])
            tpl = Path(meta["latex_path"]).read_text(encoding="utf-8")
            tex = apply_placeholders(tpl, build_placeholders(updated_json))

            build = copy_template(meta)
            extra = collect_assets(build)
            photo = updated_json.get("photo") or ""
            if photo:
                extra[photo] = app.config["UPLOAD_FOLDER"] / photo

            pdf = compile_pdf_remote(tex, build, extra)

            new_json_id = save_json(updated_json)
            new_tex_id = save_text(tex)

            session.update({
                "cv_id": new_json_id,
                "tex_id": new_tex_id,
                "pdf_filename": pdf.relative_to(app.config["OUTPUT_FOLDER"]).as_posix()
            })
            return jsonify(status="completed")
        
        elif task_id.startswith("diag_"):
             return jsonify(result)
    
    # Pour les erreurs
    elif result.get("status") == "error":
        return jsonify(status="error", message=result["message"])
    
    # Tâche encore en cours
    return jsonify(status="processing")
# Ajoutez ce dictionnaire global au début du fichier, après les autres variables




# ---------- Aperçu / téléchargement --------------------------------
@app.route("/preview")
@login_required
def preview():
    if "pdf_filename" not in session:
        return redirect(url_for("index"))

    cv_json = {}
    if "cv_id" in session:
        cv_json = load_json(session["cv_id"])

    return render_template(
        "preview.html",
        pdf_filename = session["pdf_filename"],
        ts           = uuid.uuid4().hex,           # anti-cache
        cv_json      = json.dumps(cv_json, indent=2, ensure_ascii=False)
    )


@app.route("/file/<path:filename>")
def file(filename):
#     """Servez le PDF + entêtes anti-cache."""
    resp = send_file(app.config["OUTPUT_FOLDER"] / filename, conditional=False)
    resp.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
    resp.headers["Pragma"]        = "no-cache"
    resp.headers["Expires"]       = "0"
    return resp



@app.route("/download")
@login_required 
def download():
    if "pdf_filename" not in session: return redirect(url_for("index"))
    return send_file(app.config["OUTPUT_FOLDER"]/session["pdf_filename"],
                     as_attachment=True)



def gpt_edit_json(cv_json: dict, instruction: str) -> dict:
    """
    Retourne un NOUVEAU JSON après application de *instruction*.
    * Le schéma doit rester strictement le même (mêmes clés).
    * Ne change rien d’autre que ce qui est demandé.
    * Répond UNIQUEMENT avec l’objet JSON.
    """
    sys = (
        "Tu es un assistant RH. On te donne l'objet JSON complet d'un CV "
        "et une instruction (ajouter, supprimer ou modifier). "
        "Tu appliques STRICTEMENT l'instruction, sans toucher au reste. "
        "Le résultat doit respecter exactement le même schéma et contenir "
        "toutes les clés d'origine."
    )

    user = json.dumps({
        "instruction": instruction,
        "schema": SCHEMA,
        "current_json": cv_json
    }, ensure_ascii=False, indent=2)

    rsp = openai.ChatCompletion.create(
        model           = OPENAI_MODEL,
        response_format = {"type": "json_object"},
        messages        = [
            {"role": "system", "content": sys},
            {"role": "user",   "content": user}
        ]
    )
    return json.loads(rsp.choices[0].message.content)


@app.route("/json_edit_prepare", methods=["POST"])
def json_edit_prepare():
    # le JSON courant est maintenant référencé par son identifiant
    if "cv_id" not in session:
        return redirect(url_for("preview"))

    instr = request.form["instruction"].strip()
    if not instr:
        flash("Instruction vide", "warning")
        return redirect(url_for("preview"))

    session["pending_json_edit"] = instr       # stocke l’instruction
    return render_template("edit_loading.html")  # petit spinner


@app.route("/json_edit_apply", methods=["POST"])
def json_edit_apply():
    # garde-fous
    if "pending_json_edit" not in session \
       or "cv_id"          not in session \
       or "template_id"    not in session:
        return jsonify(error="no data"), 400

    instr     = session.pop("pending_json_edit")
    original  = load_json(session["cv_id"])

    # 1) GPT  →  nouveau JSON
    updated = gpt_edit_json(original, instr)

    # 2) recompilation LaTeX + PDF
    meta  = get_template_meta(session["template_id"])
    tpl   = Path(meta["latex_path"]).read_text(encoding="utf-8")
    tex   = apply_placeholders(tpl, build_placeholders(updated))

    build = copy_template(meta)
    extra = collect_assets(build)
    if updated.get("photo"):
        extra[updated["photo"]] = app.config["UPLOAD_FOLDER"] / updated["photo"]

    pdf = compile_pdf_remote(tex, build, extra)

    # 3) sauvegarde artefacts  + MAJ session
    session["cv_id"]       = save_json(updated)
    session["tex_id"]      = save_text(tex)
    session["pdf_filename"] = pdf.relative_to(app.config["OUTPUT_FOLDER"]).as_posix()

    return jsonify(ok=True)



def gpt_edit(tex: str, instr: str) -> str:
    """Applique *instr* au code LaTeX *tex* sans rien inventer.

    * L’instruction est libre (supprimer, modifier, ajouter une ligne, etc.).
    * Aucune autre partie du document ne doit être altérée.
    * La sortie doit être le **LaTeX complet**, sans balises ``` ni commentaires.
    """

    sys = (
        "Tu es un expert en LaTeX. Tu reçois le code LaTeX complet d’un CV et une instruction unique. "
        "Applique strictement cette instruction, sans rien inventer d’autre. "
        "Ne modifie pas les parties non concernées et conserve tous les encodages/accents. "
        "Réponds UNIQUEMENT avec le LaTeX final, sans balises ``` ni commentaires supplémentaires."
    )

    user = f"INSTRUCTION : {instr}\n\nCODE LATEX :\n{tex}"

    out = openai.ChatCompletion.create(
        model       = OPENAI_MODEL,
        messages    = [
            {"role": "system", "content": sys},
            {"role": "user",   "content": user}
        ]
        
    ).choices[0].message.content

    # sécurité : retire tout bloc ``` éventuel
    return re.sub(r"^```.*?\n|\n?```$", "", out, flags=re.S).strip()

# ---------------------------------------------------------------------------
# 1️⃣  Phase rapide : on stocke l’instruction et on affiche le spinner
# ---------------------------------------------------------------------------
@app.route("/edit_prepare", methods=["POST"])
def edit_prepare():
    if "template_id" not in session or "last_tex" not in session:
        return redirect(url_for("preview"))

    instr = request.form.get("instruction", "").strip()
    if not instr:
        flash("Veuillez saisir une instruction !", "warning")
        return redirect(url_for("preview"))

    session["pending_json_edit"] = instr
    return render_template("edit_loading.html",
                       target=url_for('json_edit_apply'))

# ────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    app.run(debug=True, port=5000)